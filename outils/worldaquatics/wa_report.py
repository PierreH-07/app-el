"""Génère dans staging/ : JSON + HTML mis à jour, historique JSON, Excel avant/après et Excel historique H/F."""
import json, os, re
from collections import Counter
from datetime import date, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
import wa_process as P
import wa_history as Hh

SCR = P.SCR
ST = os.environ.get('WA_OUT', f'{SCR}/staging')
os.makedirs(ST, exist_ok=True)
FONT = 'Arial'
HF = Font(name=FONT, bold=True, color='FFFFFF', size=10)
HFILL = PatternFill('solid', fgColor='1F3864')
BF = Font(name=FONT, size=10)
TF = Font(name=FONT, bold=True, size=13, color='1F3864')
NF = Font(name=FONT, italic=True, size=9, color='666666')


def sheet(wb, title, headers, rows, widths, note=None, first=False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    r0 = 1
    if note:
        ws.cell(row=1, column=1, value=note).font = NF
        r0 = 3
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r0, column=i, value=h)
        c.font, c.fill = HF, HFILL
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for j, row in enumerate(rows, r0 + 1):
        for i, v in enumerate(row, 1):
            c = ws.cell(row=j, column=i, value=v)
            c.font = BF
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = ws.cell(row=r0 + 1, column=1)
    if rows:
        ws.auto_filter.ref = f"A{r0}:{get_column_letter(len(headers))}{r0 + len(rows)}"
    return ws


def secs(t):
    s = P.to_sec(t) if t else None
    return s


def main():
    raw, refs, n10, ko, html, log, flags = P.run()
    hist, date_log = Hh.run()

    # corrections de dates (écart > 7 j, >= 3 nageurs WA concordants)
    courses = json.load(open(f'{P.DATA}/resultats_10km_courses_el.json'))
    date_changes = []
    for d in date_log:
        src, key = d['course'].split('/', 1)
        if src == '10km':
            for g in ('CF', 'CH'):
                co = courses[g].get(key)
                if co and co.get('date') == d['date_base'] and co.get('label') == d['label']:
                    co['date'] = d['date_wa']
                    date_changes.append(dict(fichier='resultats_10km_courses_el.json', cle=f'{g}/{key}', **d))
        else:
            co = ko['KO_DATA'].get(key)
            if co and co.get('date') == d['date_base']:
                co['date'] = d['date_wa']
                date_changes.append(dict(fichier='resultats_ko_el.json', cle=f'KO_DATA/{key}', **d))

    dump = lambda o, f: open(f, 'w', encoding='utf-8').write(json.dumps(o, ensure_ascii=False, separators=(',', ':')))
    dump(n10, f'{ST}/resultats_10km_nageurs_el.json')
    dump(ko, f'{ST}/resultats_ko_el.json')
    dump(courses, f'{ST}/resultats_10km_courses_el.json')
    open(f'{ST}/historique_bassin_el.json', 'w', encoding='utf-8').write(json.dumps(hist, ensure_ascii=False, separators=(',', ':')))
    for f, txt in html.items():
        open(f'{ST}/{f}', 'w', encoding='utf-8').write(txt)

    # ================= Excel avant / après =================
    wb = Workbook()
    temps = [l for l in log if l['motif'] in ('amelioration WA', 'ajout WA (absent de la base)')]
    annees = [l for l in log if l['champ'] == 'annee']
    nett = [l for l in log if l['motif'] in ('vitesse recalculee depuis le temps', 'format normalise', 'placeholder/format invalide -> null')]
    is_fiche = lambda v: any(not x.startswith('participant') for x in v['base']['sources'])
    st = Counter(v['status'] for v in raw.values() if is_fiche(v))
    stp = Counter(v['status'] for v in raw.values() if not is_fiche(v))
    from collections import defaultdict
    byid = defaultdict(list)
    for v in raw.values():
        if v.get('status') == 'ok':
            byid[v['athlete']['id']].append(v['base']['nom'])
    for i, noms in byid.items():
        if len(noms) > 1:
            flags.append(dict(nom=' / '.join(noms), src='base', probleme=f'doublon dans la base : meme nageur WA (id {i}) sous {len(noms)} orthographes'))
    synth = [
        ['Fiches nageurs controlees (JSON + STARTLIST)', sum(st.values())],
        ['Identifies sur World Aquatics', st.get('ok', 0)],
        ['Introuvables', st.get('introuvable', 0)],
        ['Ambigus (non modifies)', st.get('ambigu', 0)],
        ['Autres erreurs', sum(v for k, v in st.items() if k not in ('ok', 'introuvable', 'ambigu'))],
        ['Participants aux courses sans fiche (historique uniquement)', sum(stp.values())],
        ['  dont identifies sur WA', stp.get('ok', 0)],
        ['', ''],
        ['Fenetre de recherche', f'{P.CUTOFF.isoformat()} -> {P.TODAY.isoformat()} (24 mois)'],
        ['Regle', 'temps WA remplace la base uniquement s il est meilleur ; bassin 50m valide ((25m) exclu)'],
        ['', ''],
        ['Temps ameliores (toutes sources)', sum(1 for l in temps if l['motif'] == 'amelioration WA')],
        ['Temps ajoutes (absents de la base)', sum(1 for l in temps if l['motif'] != 'amelioration WA')],
        ['Annees de naissance ajoutees', len(annees)],
        ['Corrections techniques (vitesses/formats)', len(nett)],
        ['Dates de course corrigees', len(date_changes)],
        ['Points a verifier', len(flags)],
    ]
    sheet(wb, 'Synthese', ['Element', 'Valeur'], synth, [42, 90], first=True)
    rows = []
    for l in temps:
        a, b = secs(l['avant']), secs(l['apres'])
        rows.append([l['src'], l['nom'], l['champ'], l['avant'], l['apres'], round(a - b, 2) if a and b else None,
                     l.get('date_perf'), l.get('competition'), l.get('validation'), l['motif']])
    sheet(wb, 'Temps modifies', ['Fichier', 'Nom', 'Champ', 'Avant', 'Apres', 'Gain (s)', 'Date perf', 'Competition', 'Validation bassin', 'Motif'],
          rows, [34, 28, 7, 11, 11, 9, 11, 48, 26, 26])
    sheet(wb, 'Annees ajoutees', ['Fichier', 'Nom', 'Avant', 'Apres'],
          [[l['src'], l['nom'], l['avant'], l['apres']] for l in annees], [34, 30, 8, 8])
    sheet(wb, 'Corrections techniques', ['Fichier', 'Nom', 'Champ', 'Avant', 'Apres', 'Motif'],
          [[l['src'], l['nom'], l['champ'], str(l['avant']) if l['avant'] is not None else None,
            str(l['apres']) if l['apres'] is not None else None, l['motif']] for l in nett], [34, 28, 7, 12, 12, 36])
    sheet(wb, 'Dates courses', ['Fichier', 'Cle', 'Course', 'Date base', 'Date WA', 'Nb nageurs concordants'],
          [[d['fichier'], d['cle'], d['label'], d['date_base'], d['date_wa'], d['nb_nageurs_concordants']] for d in date_changes],
          [30, 30, 26, 11, 11, 12])
    sheet(wb, 'A verifier', ['Nom', 'Source', 'Probleme'], [[f['nom'], f['src'], f['probleme']] for f in flags], [30, 34, 110],
          note='Non modifie automatiquement : nageur introuvable/ambigu sur WA, ou incoherence annee/NOC entre base et WA.')
    wb.save(f'{ST}/controle_WA_avant_apres.xlsx')

    # ================= Excel historique H / F =================
    for g, lab in (('F', 'Femmes'), ('H', 'Hommes')):
        wb = Workbook()
        rows = []
        for ck, c in sorted(hist['courses'].items(), key=lambda kv: (kv[1]['date'] or '', kv[0])):
            if c['genre'] != g:
                continue
            for p in c['participants']:
                if not any(p.get(str(d)) for d in P.DISTS):
                    continue
                r = [c['date'], c['label'], 'KO 3km' if c['source'] == 'KO' else c['key'], p['nom'], p.get('rang')]
                for d in P.DISTS:
                    b = p.get(str(d))
                    r += [b['t'] if b else None, b['date'] if b else None]
                r.append(p.get('vc'))
                rows.append(r)
        sheet(wb, 'Par course', ['Date course', 'Course', 'Epreuve', 'Nom', 'Rang', '400', 'Date 400', '800', 'Date 800', '1500', 'Date 1500', 'VC (m/s)'],
              rows, [11, 24, 20, 28, 6, 9, 11, 9, 11, 10, 11, 9], first=True,
              note='Meilleurs temps bassin 50m (WA) sur les 24 mois precedant la course. VC = 1100/(t1500-t400).')
        rows = []
        for k, n in sorted(hist['nageurs'].items(), key=lambda kv: kv[1]['nom']):
            if n['genre'] != g:
                continue
            for y, wbst in sorted((n.get('par_annee') or {}).items()):
                r = [n['nom'], n.get('noc'), n.get('annee'), int(y)]
                for d in P.DISTS:
                    b = wbst.get(str(d))
                    r += [b['t'] if b else None, b['date'] if b else None]
                r.append(wbst.get('vc'))
                rows.append(r)
        sheet(wb, 'Par annee', ['Nom', 'NOC', 'Annee naiss.', 'Annee', '400', 'Date 400', '800', 'Date 800', '1500', 'Date 1500', 'VC (m/s)'],
              rows, [28, 6, 8, 7, 9, 11, 9, 11, 10, 11, 9],
              note='Meilleurs temps bassin 50m (WA) sur les 24 mois se terminant au 31/12 de l annee (annee en cours : jusqu a aujourd hui).')
        rows = []
        for k, n in sorted(hist['nageurs'].items(), key=lambda kv: kv[1]['nom']):
            if n['genre'] != g:
                continue
            for p in n.get('perfs') or []:
                if p['d'] in P.DISTS:
                    rows.append([n['nom'], n.get('noc'), p['d'], p['t'], p['date'], p['bassin'], p['valid'], p['pts'], p['comp']])
        sheet(wb, 'Perfs WA', ['Nom', 'NOC', 'Dist', 'Temps', 'Date', 'Bassin retenu', 'Validation', 'Points WA', 'Competition'],
              rows, [28, 6, 6, 10, 11, 12, 30, 9, 50],
              note='Toutes les perfs 400/800/1500 NL recuperees sur WA, avec la classification bassin appliquee.')
        wb.save(f'{ST}/historique_bassin_{lab}.xlsx')
    print('staging ecrit :', sorted(os.listdir(ST)))
    print(synth)


if __name__ == '__main__':
    main()
